#!/usr/bin/env python3
# [CLAUDE-CONSTRUCTOR] Código creado por Claude como constructor del servidor. El agente de contenido no lo escribe ni lo edita.
"""Contrato JSON -> PowerPoint 16:9 EDITABLE (python-pptx). Mismo contrato, roles de color, fuentes y bloques que el PDF.

    python servidor/render_pptx.py contenido/guia.json -o salida/guia.pptx --imagenes teoria/imagenes

Principios (PLAN_PPTX.md): una idea por diapositiva; el servidor parte el contenido (≤ 6 viñetas, ≤ 8 filas por diapositiva);
texto, tablas, gráficos, diagramas y avisos son objetos nativos editables; el color de texto se elige SOLO entre los pares permitidos del manifiesto;
el tema de PowerPoint lleva la paleta y las fuentes del diseño. Lo no estándar se degrada con aviso (nunca rompe).
"""
import argparse
import sys
import json
import math
import re
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import render
from engine import assets, contract, paleta
from engine.text import plain

W, H = 13.333, 7.5            # pulgadas, 16:9
M = 0.62                      # margen lateral
CW = W - 2 * M                # ancho útil
TOP_CONT = 1.05               # inicio del cuerpo en diapositivas de continuación
BOT = 6.78                    # fin del cuerpo (debajo va el pie)
GAP = 0.22
MAX_ITEMS, MAX_ROWS, MAX_P = 6, 8, 620
FONT_SWAP = {"Anton": "Impact", "Archivo": "Arial", "IBM Plex Mono": "Consolas"}
TONES = {"nota": ("arena", "NOTA"), "clave": ("bosque", "CLAVE"), "ojo": ("ocre", "OJO"), "sim": ("crema", "SIMULACIÓN"), "pend": ("arena", "PENDIENTE"),
         "regla": ("negro", "REGLA"), "slate": ("arena", "NOTA"), "hueco": ("arena", "HUECO")}
TOK = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*(?!\s).+?(?<!\s)\*(?![\w*])|\[\^\w{1,3}\]|\[[A-ZÁÉÍÓÚ ]{2,14}\])")


CHAR_W = {"body": 0.52, "display": 0.48, "mono": 0.6}     # ancho medio de carácter en em (medido a ojo en PowerPoint)


def est_h(text, w, pt, lh=1.15, after=0.0, font="body"):
    """Alto estimado (pulgadas) de un texto en una caja de ancho w. PowerPoint: interlineado 1,0 ≈ 1,2 × el cuerpo de letra."""
    cpl = max(6, int(w * 72 / (pt * CHAR_W[font])))
    return sum(max(1, math.ceil(len(t) / cpl)) * pt * 1.2 * lh / 72 + pt * after / 72 for t in str(text).split("\n"))


def fit(text, w, h, size, minimum, lh=1.0, font="body"):
    """Regla R12 del PDF: un título que no cabe se reduce por pasos hasta el mínimo."""
    while size > minimum and est_h(text, w, size, lh, font=font) > h:
        size -= 2
    return size


def chunks(seq, n):
    return [seq[i:i + n] for i in range(0, len(seq), n)] or [[]]


def cell_txt(c):
    return plain(str(c.get("t", "") if isinstance(c, dict) else c))


def col_widths(head, rows, w, size=12, rel=None):
    """Anchos (pulgadas) que dejan en UNA línea los atributos: encabezados y primera columna (si es corta, ≤ 28 caracteres).
    El resto del ancho se reparte según el texto medio de cada columna. Si no cabe, se reduce la letra del encabezado hasta 9 pt.
    Devuelve (anchos, puntos del encabezado, sin_salto)."""
    nc = max(len(head), max((len(r) for r in rows), default=1), 1)
    if rel and len(rel) == nc:                       # anchos del autor: se respetan
        return [w * x / sum(rel) for x in rel], size - 1, False
    first = max((len(cell_txt(r[0])) for r in rows if r), default=0)
    hs = size - 1
    while True:
        mins = [len(cell_txt(head[j])) * hs * CHAR_W["mono"] * 1.12 / 72 + 0.22 if j < len(head) else 0.5 for j in range(nc)]
        if first <= 28:
            mins[0] = max(mins[0], first * size * CHAR_W["body"] * 1.05 / 72 + 0.22)
        if sum(mins) <= w or hs <= 9:
            break
        hs -= 1
    ok = sum(mins) <= w
    if not ok:
        mins = [m * w / sum(mins) for m in mins]
    avg = [max(3, sum(len(cell_txt(r[j])) for r in rows if j < len(r)) / max(len(rows), 1)) for j in range(nc)]
    extra = max(0, w - sum(mins))
    return [m + extra * a / sum(avg) for m, a in zip(mins, avg)], hs, ok


def split_p(t, n=MAX_P):
    """Parte un párrafo largo por frases (nunca a mitad de frase si se puede)."""
    if len(t) <= n:
        return [t]
    out, cur = [], ""
    for s in re.split(r"(?<=[.;:!?])\s+", t):
        if cur and len(cur) + len(s) > n:
            out.append(cur)
            cur = s
        else:
            cur = f"{cur} {s}".strip()
    return out + [cur] if cur else out


class Deck:
    def __init__(self, roles, m, images_root, report):
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = Inches(W), Inches(H)
        self.roles, self.m, self.img_root, self.rep = roles, m, Path(images_root), report
        self.C = {k: RGBColor.from_string(v.lstrip("#").upper()) for k, v in roles.items()}
        fam = [f["family"] for f in m["fonts"]]
        self.F = {"display": fam[0], "body": fam[1] if len(fam) > 1 else fam[0], "mono": fam[2] if len(fam) > 2 else fam[-1]}
        self.pairs = {tuple(x) for x in m["palette"]["pares_texto_permitidos"]}
        self.credits = assets.load_credits(self.img_root)
        self.used_credits, self.cache = {}, Path(tempfile.gettempdir()) / "anura-pptx-cache"
        self.footer, self.n = "", 0
        self._theme()

    # ---------- tema: la paleta y las fuentes del diseño viajan en el .pptx (Diseño > Variantes las muestra) ----------
    def _theme(self):
        th = self.p.slide_master.part.part_related_by(RT.THEME)
        x = etree.fromstring(th.blob)
        a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        slots = {"dk1": "negro", "lt1": "crema", "dk2": "bosque", "lt2": "arena", "accent1": "hoja", "accent2": "vivo", "accent3": "ocre", "accent4": "bosque",
                 "accent5": "negro", "accent6": "arena", "hlink": "hoja", "folHlink": "bosque"}
        cs = x.find(f".//{a}clrScheme")
        cs.set("name", self.m["name"])
        for slot, role in slots.items():
            el = cs.find(f"{a}{slot}")
            for c in list(el):
                el.remove(c)
            etree.SubElement(el, f"{a}srgbClr", val=self.roles[role].lstrip("#").upper())
        fs = x.find(f".//{a}fontScheme")
        fs.set("name", self.m["name"])
        fs.find(f"{a}majorFont/{a}latin").set("typeface", self.F["display"])
        fs.find(f"{a}minorFont/{a}latin").set("typeface", self.F["body"])
        th._blob = etree.tostring(x, xml_declaration=True, encoding="UTF-8", standalone=True)

    # ---------- color de texto: solo pares permitidos (texto, fondo) del manifiesto ----------
    def fg(self, bg, prefs=("negro", "crema", "bosque")):
        for r in prefs:
            if (r, bg) in self.pairs:
                return r
        for r in ("negro", "crema"):
            if paleta.ratio(self.roles[r], self.roles[bg]) >= 4.5:
                return r
        return "negro"

    def accent(self, bg):
        return self.fg(bg, ("vivo", "ocre", "crema", "negro"))

    # ---------- primitivas ----------
    def slide(self, bg="crema"):
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        f = s.background.fill
        f.solid()
        f.fore_color.rgb = self.C[bg]
        s._bg = bg
        self.n += 1
        return s

    def rect(self, s, x, y, w, h, fill=None, line="negro", lw=2.25, shape=MSO_SHAPE.RECTANGLE, dash=False):
        r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        r.shadow.inherit = False
        if fill:
            r.fill.solid()
            r.fill.fore_color.rgb = self.C[fill]
        else:
            r.fill.background()
        if line:
            r.line.color.rgb = self.C[line]
            r.line.width = Pt(lw)
            if dash:
                r.line.dash_style = 4    # MSO_LINE.DASH
        else:
            r.line.fill.background()
        return r

    def runs(self, par, t, size, color, font="body", bold=False, italic=False):
        for piece in TOK.split(str(t)):
            if not piece:
                continue
            r = par.add_run()
            f, b, i, sz, txt = font, bold, italic, size, piece
            if piece.startswith("**") and piece.endswith("**"):
                txt, b = piece[2:-2], True
            elif piece.startswith("`") and piece.endswith("`"):
                txt, f, sz = piece[1:-1], "mono", size * 0.9
            elif piece.startswith("[^"):
                txt, sz = f"[{piece[2:-1]}]", size * 0.7
            elif re.fullmatch(r"\[[A-ZÁÉÍÓÚ ]{2,14}\]", piece):
                txt, f, sz = piece, "mono", size * 0.68
            elif len(piece) > 2 and piece[0] == "*" and piece[-1] == "*":
                txt, i = piece[1:-1], True
            r.text = txt
            r.font.name, r.font.size, r.font.bold, r.font.italic = self.F[f], Pt(round(sz, 1)), b, i
            r.font.color.rgb = self.C[color]

    def tb(self, s, x, y, w, h, paras, size, color, font="body", bold=False, align=None, anchor=MSO_ANCHOR.TOP, lh=1.15, after=0.35, upper=False, check=True):
        paras = [paras] if isinstance(paras, (str, int, float)) else paras
        sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = sh.text_frame
        tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, t in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing, p.space_after = lh, Pt(size * after)
            if align is not None:
                p.alignment = align
            self.runs(p, str(t).upper() if upper else t, size, color, font, bold)
        if check:
            need = sum(est_h(plain(t).upper() if upper else plain(t), w, size, lh, after, font) for t in paras)
            if need > h * 1.12 + 0.05:
                self.rep["auditoria"].setdefault("desbordes", []).append({"diapositiva": self.n, "texto": plain(paras[0])[:60], "alto": round(need, 2), "caja": round(h, 2)})
        return sh

    def bullets(self, s, x, y, w, h, items, size, color, ordered=False, start=1):
        sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = sh.text_frame
        tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        first = True
        for it in items:
            for lvl, t in ([(0, it)] if not isinstance(it, dict) else [(0, it.get("t", ""))] + [(1, sub) for sub in it.get("items", [])]):
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.line_spacing, p.space_after = 1.12, Pt(size * 0.45)
                self.runs(p, t, size - 2 * lvl, color)
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(Inches(0.34 + 0.34 * lvl)))
                pPr.set("indent", str(-Inches(0.34)))
                etree.SubElement(pPr, qn("a:buClr")).append(etree.Element(qn("a:srgbClr"), val=self.roles[self.accent(s._bg) if lvl == 0 else color].lstrip("#")))
                if ordered and lvl == 0:
                    etree.SubElement(pPr, qn("a:buFont"), typeface=self.F["mono"])
                    etree.SubElement(pPr, qn("a:buAutoNum"), type="arabicPeriod", startAt=str(start))
                else:
                    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
                    etree.SubElement(pPr, qn("a:buChar"), char="■" if lvl == 0 else "–")
        return sh

    def notes(self, s, text):
        if text:
            tf = s.notes_slide.notes_text_frame
            tf.text = (tf.text + "\n" + text).strip()

    def chrome(self, s, crumb):
        """Cabecera (miga de pan) y pie con número: iguales en todas las hojas de contenido."""
        ink = self.fg(s._bg)
        self.tb(s, M, 0.38, CW * 0.7, 0.3, crumb, 11, ink, "mono", True, upper=True, check=False)
        self.rect(s, M, 0.74, CW, 0, line=ink, lw=3)
        self.tb(s, M, 7.0, CW * 0.8, 0.25, self.footer, 10, ink, "mono", check=False)
        self.tb(s, W - M - 1, 7.0, 1, 0.25, f"{self.n:02d}", 10, ink, "mono", True, align=PP_ALIGN.RIGHT, check=False)

    # ---------- imágenes ----------
    def picture(self, s, src, x, y, w, h, treatment="calido", need=None, caption=None, decor=False):
        f = self.img_root / src if src else None
        if not f or not f.exists():
            if decor:      # portada y divisores: sin la foto la composición sigue funcionando; solo se informa
                self.rep["imagenes_necesarias"].append({"diapositiva": self.n, "imagen": src or "", "que": plain(need or caption or "foto decorativa")})
                return False
            self.rect(s, x, y, w, h, "arena", dash=True)
            self.tb(s, x + 0.2, y + h / 2 - 0.3, w - 0.4, 0.6, ["IMAGEN NECESARIA", need or src or ""], 12, self.fg("arena"), "mono", True, align=PP_ALIGN.CENTER, check=False)
            self.rep["imagenes_necesarias"].append({"diapositiva": self.n, "imagen": src or "", "que": plain(need or caption or "")})
            return
        try:
            f = assets.duotone(f, treatment or "calido", self.roles, self.m["image_treatments"], self.cache)
        except Exception as e:  # noqa: BLE001
            self.rep["avisos"].append({"bloque": src, "causa": f"no se pudo aplicar el filtro: {e}", "ajuste": "se usó la imagen original"})
        pic = s.shapes.add_picture(str(f), Inches(x), Inches(y))
        r = min(Inches(w) / pic.width, Inches(h) / pic.height)
        pic.width, pic.height = int(pic.width * r), int(pic.height * r)
        pic.left, pic.top = int(Inches(x) + (Inches(w) - pic.width) / 2), int(Inches(y) + (Inches(h) - pic.height) / 2)
        pic.line.color.rgb, pic.line.width = self.C["negro"], Pt(2.25)
        if caption:
            pic.element.nvPicPr.cNvPr.set("descr", plain(caption))
        c = self.credits.get(src) or self.credits.get(Path(src).name)
        if c:
            self.used_credits[src] = c
        return True

    # ---------- gráficos nativos ----------
    def chart(self, s, ch, x, y, w, h, b):
        t, d = ch.get("type"), ch.get("data") or {}
        SER = ["hoja", "vivo", "negro", "ocre", "arena", "bosque", "crema"]
        ink = self.fg(s._bg)
        multi, kind = False, None
        if t in ("bar", "hbar", "hist", "line", "stack", "pie"):
            cd = CategoryChartData()
            if t in ("bar", "hbar") and d.get("series") and d.get("categories"):
                cd.categories = [str(c) for c in d["categories"]]
                for sr in d["series"]:
                    cd.add_series(sr.get("name", ""), sr.get("values", []))
                multi = True
                kind = XL_CHART_TYPE.COLUMN_STACKED if d.get("mode") == "stacked" else XL_CHART_TYPE.COLUMN_CLUSTERED
            elif t in ("bar", "hbar", "pie"):
                its = d.get("items", [])
                cd.categories = [plain(i.get("l", "")) for i in its]
                cd.add_series(plain(ch.get("title") or b.get("title") or "Serie"), [i.get("v", 0) for i in its])
                kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "hbar": XL_CHART_TYPE.BAR_CLUSTERED, "pie": XL_CHART_TYPE.DOUGHNUT if d.get("donut") else XL_CHART_TYPE.PIE}[t]
            elif t == "hist":
                cd.categories = [f"{x_:g}".replace(".", ",") for x_, _ in d.get("bars", [])]
                cd.add_series(d.get("ylabel", "Conteo"), [v for _, v in d.get("bars", [])])
                kind = XL_CHART_TYPE.COLUMN_CLUSTERED
            elif t == "line":
                series = d.get("series") or [{"name": d.get("ylabel", ""), "points": d.get("points", [])}]
                xs = list(dict.fromkeys(p[0] for sr in series for p in sr["points"]))
                cd.categories = [str(v) for v in xs]
                for sr in series:
                    look = {p[0]: p[1] for p in sr["points"]}
                    cd.add_series(sr.get("name", ""), [look.get(v) for v in xs])
                multi, kind = len(series) > 1, XL_CHART_TYPE.LINE_MARKERS
            else:   # stack (composición 100 %)
                names = d.get("series") or [f"Parte {i+1}" for i in range(len(d["rows"][0]["v"]))]
                cd.categories = [plain(r["l"]) for r in d.get("rows", [])]
                for i, nm in enumerate(names):
                    cd.add_series(str(nm), [r["v"][i] if i < len(r["v"]) else 0 for r in d["rows"]])
                multi, kind = True, XL_CHART_TYPE.BAR_STACKED_100
        elif t == "scatter":
            cd = XyChartData()
            sr = cd.add_series(d.get("ylabel", ""))
            for px, py in (p[:2] for p in d.get("points", [])):
                sr.add_data_point(px, py)
            kind = XL_CHART_TYPE.XY_SCATTER
        else:
            return self.degrade_chart(s, t, d, x, y, w, h)
        try:
            gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        except (ValueError, KeyError, TypeError, IndexError) as e:      # datos vacíos o mal formados: nunca rompe
            self.rep["avisos"].append({"bloque": f"diapositiva {self.n}", "causa": f"gráfico '{t}' sin datos válidos ({e})", "ajuste": "revisa `chart.data`"})
            return self.rect(s, x, y, w, h, "arena", dash=True)
        c = gf.chart
        c.font.name, c.font.size = self.F["body"], Pt(12)
        c.font.color.rgb = self.C[ink]
        c.has_title = False
        c.has_legend = multi or t == "pie"
        if c.has_legend:
            c.legend.position, c.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
            c.legend.font.size = Pt(11)
        plot = c.plots[0]
        if t in ("bar", "hbar", "hist", "stack"):
            plot.gap_width = 4 if t == "hist" else 60
            if multi and t != "stack" and d.get("mode") != "stacked":
                plot.overlap = -10
        for j, ser in enumerate(plot.series):
            fmt = ser.format
            if t == "line":
                fmt.line.color.rgb, fmt.line.width = self.C[SER[j % 5] if multi else "hoja"], Pt(3)
                ser.marker.style, ser.marker.size = XL_MARKER_STYLE.SQUARE, 8
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = self.C["negro"]
                ser.smooth = False
            elif t == "scatter":
                ser.marker.style, ser.marker.size = XL_MARKER_STYLE.CIRCLE, 8
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = self.C["hoja"]
                ser.marker.format.line.color.rgb = self.C["negro"]
                fmt.line.fill.background()
            elif t != "pie":
                fmt.fill.solid()
                fmt.fill.fore_color.rgb = self.C[SER[j % 5] if multi else "hoja"]
                fmt.line.color.rgb, fmt.line.width = self.C["negro"], Pt(0.75)
        if t in ("bar", "hbar", "pie") and not multi:
            for k, it in enumerate(d.get("items", [])):
                pt_ = plot.series[0].points[k]
                pt_.format.fill.solid()
                pt_.format.fill.fore_color.rgb = self.C[SER[k % len(SER)] if t == "pie" else ("vivo" if it.get("hi") else "hoja")]
                pt_.format.line.color.rgb = self.C["negro"]
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size, dl.font.bold = Pt(11), True
            dl.font.color.rgb = self.C[ink]
            if t != "pie":
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
        if t == "scatter" and d.get("trend"):
            ser = plot.series[0]._element
            ser.find(qn("c:xVal")).addprevious(etree.fromstring(
                f'<c:trendline xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                f'<c:spPr><a:ln w="28575"><a:solidFill><a:srgbClr val="{self.roles["negro"].lstrip("#")}"/></a:solidFill><a:prstDash val="dash"/></a:ln></c:spPr>'
                f'<c:trendlineType val="linear"/></c:trendline>'))
        if t != "pie":
            va = c.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = self.C["arena" if s._bg != "arena" else "crema"]
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(11)
            ca = c.category_axis
            ca.format.line.color.rgb = self.C[ink]
            ca.tick_labels.font.size = Pt(11)
            if t == "hbar":
                ca.reverse_order = True
            rng = d.get("yrange") or ([0, d["max"]] if d.get("max") else None)
            if rng and t != "stack":
                va.minimum_scale, va.maximum_scale = rng
            if t == "scatter" and d.get("xrange"):
                ca.minimum_scale, ca.maximum_scale = d["xrange"]
            for ax, lab in ((ca, d.get("xlabel")), (va, d.get("ylabel"))):
                if lab:
                    ax.has_title = True
                    ax.axis_title.text_frame.text = plain(lab)
                    ax.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
        return gf

    def degrade_chart(self, s, t, d, x, y, w, h):
        """box, heat y map no tienen equivalente nativo editable: se dibujan como tabla nativa (heat con celdas coloreadas)."""
        self.rep["avisos"].append({"bloque": f"diapositiva {self.n}", "causa": f"gráfico '{t}' sin equivalente nativo en PowerPoint", "ajuste": "se muestra como tabla editable"})
        if t == "heat":
            lab, mx = d.get("labels", []), d.get("matrix", [])
            lo, hi = (d.get("range") or [-1, 1])
            head, rows = [""] + [plain(v) for v in lab], [[plain(lab[i]) if i < len(lab) else ""] + [f"{v:g}".replace(".", ",") for v in r] for i, r in enumerate(mx)]
            tbl = self.table(s, head, rows, x, y, w, h, size=10)
            for i, r in enumerate(mx):
                for j, v in enumerate(r):
                    k = (v - lo) / ((hi - lo) or 1)
                    role = "bosque" if k > .85 else "hoja" if k > .65 else "arena" if k > .35 else "crema"
                    cell = tbl.cell(i + 1, j + 1)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self.C[role]
                    for p in cell.text_frame.paragraphs:
                        for r_ in p.runs:
                            r_.font.color.rgb = self.C[self.fg(role)]
            return tbl
        if t == "box":
            head = ["Grupo", "Mín.", "Q1", "Mediana", "Q3", "Máx.", "Atípicos"]
            rows = [[plain(g.get("l") or "—")] + [f"{g.get(k, '')}".replace(".", ",") for k in ("lo", "q1", "med", "q3", "hi")] + [", ".join(f"{v:g}" for v in g.get("out", []))] for g in d.get("groups", [])]
        elif t == "map":
            head = ["Lugar", "Longitud", "Latitud", "Valor"]
            rows = [[plain(p.get("label", "")), f"{p.get('lon')}", f"{p.get('lat')}", f"{p.get('v', '')}"] for p in d.get("points", [])]
        else:
            head, rows = ["Dato"], [[json.dumps(d, ensure_ascii=False)[:200]]]
        return self.table(s, head, rows, x, y, w, min(h, 0.42 * (len(rows) + 1)))

    def table(self, s, head, rows, x, y, w, h, widths=None, size=12):
        nr, nc = len(rows) + (1 if head else 0), max(len(head or []), max((len(r) for r in rows), default=1))
        gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h))
        tbl = gf.table
        tbl.first_row = bool(head)
        ws, hs, one_line = col_widths(head or [], rows, w, size, widths)
        first_short = max((len(cell_txt(r[0])) for r in rows if r), default=0) <= 28
        for j, wj in enumerate(ws):
            tbl.columns[j].width = Inches(wj)
        cells = ([head] if head else []) + rows
        for i, r in enumerate(cells):
            hd = head and i == 0
            for j in range(nc):
                v = r[j] if j < len(r) else ""
                hi = isinstance(v, dict) and v.get("hi")
                txt = v.get("t", "") if isinstance(v, dict) else v
                bg = "negro" if hd else ("ocre" if hi else ("arena" if i % 2 == 0 else "crema"))
                cell = tbl.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.C[bg]
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.04)
                tf = cell.text_frame
                tf.word_wrap = not (one_line and (hd or (j == 0 and first_short)))     # atributos en una línea
                p = tf.paragraphs[0]
                p.text = ""
                self.runs(p, str(txt).upper() if hd else (f"{txt:g}".replace(".", ",") if isinstance(txt, float) else txt),
                          hs if hd else size, self.fg(bg), "mono" if hd else "body", bool(hd or hi))
        tbl.rows[0].height = Inches(0.36)
        return tbl

    # ---------- unidades de contenido: (alto, dibujar) ----------
    def units(self, blocks):
        """Convierte bloques en unidades medibles; parte lo largo (listas, tablas, párrafos) y empareja `half`."""
        out, i = [], 0
        while i < len(blocks):
            b = blocks[i]
            if b.get("half") and i + 1 < len(blocks) and blocks[i + 1].get("half"):
                a, c = self.unit(b, CW / 2 - GAP / 2), self.unit(blocks[i + 1], CW / 2 - GAP / 2)
                out.append((max(a[0][0], c[0][0]), lambda s, x, y, w, h, a=a[0], c=c[0]: (a[1](s, x, y, w / 2 - GAP / 2, h), c[1](s, x + w / 2 + GAP / 2, y, w / 2 - GAP / 2, h))))
                out += a[1:] + c[1:]
                i += 2
                continue
            out += self.unit(b, CW)
            i += 1
        return out

    def unit(self, b, w):
        t = b.get("type")
        fn = getattr(self, f"u_{t}", None)
        if not fn:
            self.rep["avisos"].append({"bloque": t or "?", "causa": f"bloque '{t}' fuera del vocabulario", "ajuste": "se mostró como nota"})
            return self.u_callout({"tone": "nota", "text": plain(b.get("text") or b.get("html") or json.dumps(b, ensure_ascii=False))[:400]}, w)
        return fn(b, w)

    def framed(self, b, w, h_body, draw):
        """Título + cuerpo + pie, compartido por gráficos, figuras, imágenes, diagramas."""
        th = 0.5 if b.get("title") else 0
        ch = 0.42 if b.get("caption") else 0
        h = th + h_body + ch

        def d(s, x, y, w_, h_):
            ink = self.fg(s._bg)
            if th:
                self.tb(s, x, y, w_ * 0.75, 0.42, b["title"], 16, ink, bold=True, check=False)
            if b.get("tag"):
                self.tb(s, x + w_ * 0.75, y + 0.05, w_ * 0.25, 0.35, b["tag"], 10, ink, "mono", True, align=PP_ALIGN.RIGHT, upper=True, check=False)
            draw(s, x, y + th, w_, h_ - th - ch)
            if ch:
                self.tb(s, x, y + h_ - ch + 0.06, w_, ch, b["caption"], 11, ink, check=False)
        return [(h, d)]

    def u_p(self, b, w):
        return [(est_h(plain(t), w, 18, 1.12) + 0.05, lambda s, x, y, w_, h, t=t: self.tb(s, x, y, w_, h + 0.1, t, 18, self.fg(s._bg), lh=1.12, after=0))
                for t in split_p(b.get("text", ""))]

    def u_h(self, b, w):
        def d(s, x, y, w_, h):
            ink = self.fg(s._bg)
            self.tb(s, x, y, w_, 0.32, b.get("text", ""), 12, ink, "mono", True, upper=True, check=False)
            self.rect(s, x, y + 0.36, w_, 0, line=ink, lw=1.5)
        return [(0.46, d)]

    def u_list(self, b, w):
        out, start = [], 1
        for ch in chunks(b.get("items", []), MAX_ITEMS):
            txt = [it.get("t", "") if isinstance(it, dict) else it for it in ch] + [sub for it in ch if isinstance(it, dict) for sub in it.get("items", [])]
            h = sum(est_h(plain(t), w - 0.75, 18, 1.12, 0.45) for t in txt) + 0.05
            out.append((h, lambda s, x, y, w_, h_, ch=ch, st=start: self.bullets(s, x, y, w_, h_ + 0.1, ch, 18, self.fg(s._bg), b.get("ordered"), st)))
            start += len(ch)
        return out

    def u_table(self, b, w):
        head, rows, out = b.get("head") or [], b.get("rows") or [], []
        ws = col_widths(head, rows, w, 12, b.get("widths"))[0]
        rh = lambda r: max([0.36] + [est_h(cell_txt(c), ws[j] - 0.16, 12, 1.0) + 0.12 for j, c in enumerate(r) if j < len(ws)])
        cur, hh = [], 0
        groups = []
        for r in rows:
            if cur and (len(cur) >= MAX_ROWS or hh + rh(r) > 4.6):
                groups.append(cur)
                cur, hh = [], 0
            cur.append(r)
            hh += rh(r)
        groups.append(cur)
        for g in groups:
            hb = 0.38 + sum(rh(r) for r in g)
            last = g is groups[-1]
            bb = {"caption": b.get("caption")} if last else {}
            out += self.framed(bb, w, hb, lambda s, x, y, w_, h_, g=g: self.table(s, head, g, x, y, w_, h_, b.get("widths")))
        return out

    def u_chart(self, b, w):
        hb = 3.5 if b.get("half") else 4.1
        d = (b.get("chart") or {}).get("data") or {}
        n_rows = {"box": len(d.get("groups", [])), "map": len(d.get("points", []))}.get((b.get("chart") or {}).get("type"))
        if n_rows is not None:          # se degrada a tabla: alto de la tabla, no del gráfico
            hb = min(hb, 0.42 * (n_rows + 1) + 0.1)
        if b.get("sim"):
            self.rep["avisos"].append({"bloque": plain(b.get("title", "gráfico")), "causa": "datos simulados", "ajuste": "marco punteado en la diapositiva"})

        def d(s, x, y, w_, h_):
            if b.get("sim"):
                self.rect(s, x - 0.06, y - 0.04, w_ + 0.12, h_ + 0.08, line=self.fg(s._bg), dash=True, lw=1.5)
            self.chart(s, b.get("chart") or {}, x, y, w_, h_, b)
        return self.framed(b, w, hb, d)

    def u_figure(self, b, w):
        side = b.get("side") or []

        def d(s, x, y, w_, h_):
            wl = w_ * 0.68
            if b.get("chart"):
                self.chart(s, b["chart"], x, y, wl, h_, b)
            else:
                im = b.get("image") or {}
                self.picture(s, im.get("src"), x, y, wl, h_, im.get("treatment"), im.get("need"), im.get("caption"))
            xs, ws = x + wl + 0.35, w_ - wl - 0.35
            self.rect(s, xs - 0.18, y, 0, h_, line=self.fg(s._bg), lw=3)
            yy = y
            if b.get("side_title"):
                self.tb(s, xs, yy, ws, 0.3, b["side_title"], 11, self.fg(s._bg), "mono", True, upper=True, check=False)
                yy += 0.4
            step = min(1.05, (h_ - (yy - y) - (0.6 if b.get("side_note") else 0)) / max(len(side), 1))
            for it in side:
                self.tb(s, xs, yy, ws, step * 0.62, str(it.get("v", "")), 34, self.accent(s._bg) if s._bg != "crema" else "hoja", "display", check=False)
                self.tb(s, xs, yy + step * 0.6, ws, step * 0.4, it.get("l", ""), 12, self.fg(s._bg), check=False)
                yy += step
            if b.get("side_note"):
                self.tb(s, xs, y + h_ - 0.55, ws, 0.55, b["side_note"], 11, self.fg(s._bg), check=False)
        return self.framed(b, w, 4.1, d)

    def u_kpi(self, b, w):
        its = b.get("items", [])[:4]

        def d(s, x, y, w_, h_):
            n = max(len(its), 1)
            cw = (w_ - GAP * (n - 1)) / n
            for k, it in enumerate(its):
                cx = x + k * (cw + GAP)
                bg = "negro" if k == 0 else "bosque"
                self.rect(s, cx, y, cw, h_, bg)
                self.tb(s, cx + 0.2, y + 0.15, cw - 0.4, 0.8, str(it.get("v", "")), 40, self.accent(bg), "display", check=False)
                self.tb(s, cx + 0.2, y + 0.95, cw - 0.4, 0.45, it.get("l", ""), 13, self.fg(bg, ("crema",)), bold=True, check=False)
                if it.get("note"):
                    self.tb(s, cx + 0.2, y + h_ - 0.32, cw - 0.4, 0.25, it["note"], 10, self.fg(bg, ("crema",)), "mono", check=False)
        return [(1.75, d)]

    def u_stats(self, b, w):
        its = b.get("items", [])

        def d(s, x, y, w_, h_):
            ink = self.fg(s._bg)
            yy = y
            if b.get("title"):
                self.tb(s, x, yy, w_, 0.3, b["title"], 11, ink, "mono", True, upper=True, check=False)
                yy += 0.4
            n = max(len(its), 1)
            cw = w_ / n
            for k, it in enumerate(its):
                cx = x + k * cw
                if k:
                    self.rect(s, cx - 0.1, yy, 0, 1.05, line=ink, lw=3)
                self.tb(s, cx + 0.05, yy, cw - 0.2, 0.7, str(it.get("v", "")), 36, ink, "display", check=False)
                self.tb(s, cx + 0.05, yy + 0.72, cw - 0.2, 0.35, it.get("l", ""), 12, ink, check=False)
            if b.get("note"):
                self.tb(s, x, yy + 1.15, w_, 0.3, b["note"], 11, ink, check=False)
        return [(1.15 + (0.4 if b.get("title") else 0) + (0.35 if b.get("note") else 0), d)]

    def u_big(self, b, w):
        def d(s, x, y, w_, h_):
            ink = self.fg(s._bg, ("bosque", "negro", "crema"))
            self.tb(s, x, y, w_ * 0.62, h_, f"{b.get('value', '')}{b.get('unit', '')}", 150, ink, "display", anchor=MSO_ANCHOR.MIDDLE, lh=0.9, check=False)
            self.rect(s, x + w_ * 0.64, y + 0.4, 0, h_ - 0.8, line=self.fg(s._bg), lw=4)
            self.tb(s, x + w_ * 0.67, y, w_ * 0.33, h_, b.get("caption", ""), 22, self.fg(s._bg), bold=True, anchor=MSO_ANCHOR.MIDDLE)
        return [(2.7, d)]

    def u_callout(self, b, w):
        bg, lab = TONES.get(b.get("tone", "nota"), TONES["nota"])
        paras = ([f"**{b['title']}**"] if b.get("title") else []) + (b.get("paras") or [b.get("text", "")])
        h = 0.55 + sum(est_h(plain(p), w - 0.6, 16, 1.15, 0.3) for p in paras)

        def d(s, x, y, w_, h_):
            self.rect(s, x, y, w_, h_, bg, dash=b.get("tone") == "sim", lw=3)
            self.rect(s, x + 0.25, y - 0.16, 0.25 + 0.11 * len(b.get("label") or lab), 0.32, "negro", lw=0)
            self.tb(s, x + 0.33, y - 0.13, 3, 0.28, b.get("label") or lab, 10, self.fg("negro", ("crema",)), "mono", True, upper=True, check=False)
            self.tb(s, x + 0.3, y + 0.3, w_ - 0.6, h_ - 0.4, paras, 16, self.fg(bg), lh=1.15, after=0.3)
        return [(h, d)]

    def u_quote(self, b, w):
        h = est_h(plain(b.get("text", "")), w - 1.2, 26, 1.15) + (0.6 if b.get("by") else 0.2)

        def d(s, x, y, w_, h_):
            ink = self.fg(s._bg)
            self.tb(s, x, y - 0.25, 1, 1.2, "“", 110, self.accent(s._bg) if s._bg != "crema" else "hoja", "display", check=False)
            self.tb(s, x + 1.0, y, w_ - 1.2, h_ - 0.5, b.get("text", ""), 26, ink, bold=True, lh=1.15)
            if b.get("by"):
                self.tb(s, x + 1.0, y + h_ - 0.4, w_ - 1.2, 0.3, "— " + str(b["by"]), 12, ink, "mono", True, upper=True, check=False)
        return [(h, d)]

    def u_code(self, b, w):
        lines = b.get("code", "").split("\n")[:22]
        h = 0.5 + len(lines) * 0.24

        def d(s, x, y, w_, h_):
            self.rect(s, x, y, w_, h_, "negro")
            self.tb(s, x + w_ - 2.2, y + 0.1, 2, 0.25, b.get("lang", ""), 10, self.accent("negro"), "mono", True, align=PP_ALIGN.RIGHT, upper=True, check=False)
            tb = self.tb(s, x + 0.25, y + 0.25, w_ - 0.5, h_ - 0.35, "", 12, self.fg("negro", ("crema",)), "mono", check=False)
            tb.text_frame.paragraphs[0].text = ""
            for k, ln in enumerate(lines):     # el código no se interpreta como marcas
                p = tb.text_frame.paragraphs[0] if k == 0 else tb.text_frame.add_paragraph()
                r = p.add_run()
                r.text = ln or " "
                r.font.name, r.font.size = self.F["mono"], Pt(12)
                r.font.color.rgb = self.C[self.fg("negro", ("crema",))]
        return self.framed({"caption": b.get("caption")}, w, h, d)

    def u_refs(self, b, w):
        out = []
        for k, ch in enumerate(chunks(b.get("items", []), 8)):
            h = 0.45 + sum(est_h(plain(t), w - 0.75, 13, 1.12, 0.45) for t in ch)
            out.append((h, lambda s, x, y, w_, h_, ch=ch, k=k: (
                self.tb(s, x, y, w_, 0.3, b.get("title") or "Referencias", 12, self.fg(s._bg), "mono", True, upper=True, check=False),
                self.bullets(s, x, y + 0.42, w_, h_ - 0.4, ch, 13, self.fg(s._bg), True, k * 8 + 1))))
        return out

    def u_form(self, b, w):
        rows = b.get("rows", [])

        def d(s, x, y, w_, h_):
            ink = self.fg(s._bg)
            yy = y
            if b.get("title"):
                self.tb(s, x, yy, w_, 0.35, b["title"], 16, ink, bold=True, check=False)
                yy += 0.5
            for r in rows:
                self.tb(s, x, yy + 0.08, w_ * 0.3, 0.3, r, 13, ink, "mono", True, upper=True, check=False)
                self.rect(s, x + w_ * 0.32, yy + 0.4, w_ * 0.68, 0, line=ink, lw=1.5)
                yy += 0.52
        return [((0.5 if b.get("title") else 0) + 0.52 * len(rows), d)]

    def u_flow(self, b, w):
        steps = b.get("steps", [])
        cols = int(b.get("cols") or min(len(steps), 4) or 1)
        rows = math.ceil(len(steps) / cols)

        def d(s, x, y, w_, h_):
            cw, rh = (w_ - GAP * 1.6 * (cols - 1)) / cols, (h_ - GAP * (rows - 1)) / rows
            for k, st in enumerate(steps):
                r, c = divmod(k, cols)
                cx, cy = x + c * (cw + GAP * 1.6), y + r * (rh + GAP)
                bg = "vivo" if st.get("on") else "crema"
                self.rect(s, cx, cy, cw, rh, bg, lw=3)
                self.rect(s, cx, cy, 0.55, 0.5, "negro", lw=0)
                self.tb(s, cx, cy + 0.04, 0.55, 0.45, str(st.get("n") or k + 1), 22, self.accent("negro"), "display", align=PP_ALIGN.CENTER, check=False)
                self.tb(s, cx + 0.7, cy + 0.1, cw - 0.85, 0.4, st.get("title", ""), 15, self.fg(bg), bold=True, check=False)
                if st.get("text"):
                    self.tb(s, cx + 0.18, cy + 0.62, cw - 0.36, rh - 0.7, st["text"], 12, self.fg(bg))
                if c < cols - 1 and k < len(steps) - 1:
                    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + cw + 0.04), Inches(cy + rh / 2 - 0.12), Inches(GAP * 1.6 - 0.08), Inches(0.24))
                    a.fill.solid()
                    a.fill.fore_color.rgb = self.C[self.fg(s._bg)]
                    a.line.fill.background()
        txt_h = max([est_h(plain(st.get("text", "")), (CW / cols) - 0.4, 12, 1.15, 0.35) for st in steps] or [0])
        return [(rows * (0.85 + txt_h) + GAP * (rows - 1), d)]

    def u_ladder(self, b, w):
        its = b.get("items", [])

        def d(s, x, y, w_, h_):
            rh = (h_ - 0.08 * (len(its) - 1)) / max(len(its), 1)
            for k, it in enumerate(its):
                off = k * 0.35
                bg = "bosque" if k == len(its) - 1 else ("arena" if k % 2 else "crema")
                self.rect(s, x + off, y + k * (rh + 0.08), w_ - off, rh, bg)
                ww, tw = (w_ - off) * 0.3, (w_ - off) * 0.48
                self.tb(s, x + off + 0.2, y + k * (rh + 0.08) + 0.04, ww, rh - 0.08, it.get("title", ""), fit(plain(it.get("title", "")), ww, rh - 0.1, 14, 10), self.fg(bg, ("negro", "crema")), bold=True, anchor=MSO_ANCHOR.MIDDLE)
                self.tb(s, x + off + (w_ - off) * 0.32, y + k * (rh + 0.08) + 0.04, tw, rh - 0.08, it.get("text", ""), fit(plain(it.get("text", "")), tw, rh - 0.1, 12, 9.5, 1.15), self.fg(bg, ("negro", "crema")), anchor=MSO_ANCHOR.MIDDLE)
                if it.get("cost"):
                    self.tb(s, x + w_ - 1.9, y + k * (rh + 0.08) + 0.08, 1.75, rh - 0.1, it["cost"], 10, self.fg(bg, ("negro", "crema")), "mono", True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, upper=True, check=False)
        return [(min(5.2, max(0.62, 0.3 + max((est_h(plain(i.get("text", "")), CW * 0.4, 12, 1.15) for i in its), default=0)) * len(its)), d)]

    def u_ranks(self, b, w):
        rows = [[r.get("rank", "")] + list(r.get("v", [])) for r in b.get("rows", [])]
        return self.u_table({"head": ["Rango"] + list(b.get("head") or []), "rows": rows}, w)

    def u_timeline(self, b, w):
        its, ad = b.get("items", []), b.get("adult")
        vis = {"v": "vivo", "h": "ocre", "x": "arena"}

        def d(s, x, y, w_, h_):
            tot = sum(float(it.get("days") or 1) for it in its) or 1
            wa = 1.6 if ad else 0
            cx = x
            for it in its:
                cw = (w_ - wa) * float(it.get("days") or 1) / tot
                bg = vis.get(it.get("vis"), "arena")
                self.rect(s, cx, y, cw, 1.0, bg)
                self.tb(s, cx + 0.12, y + 0.08, cw - 0.2, 0.6, it.get("code", ""), 30, self.fg(bg), "display", check=False)
                self.tb(s, cx + 0.12, y + 0.66, cw - 0.2, 0.3, f"{it.get('days', '')} d", 10, self.fg(bg), "mono", True, check=False)
                self.tb(s, cx, y + 1.1, cw - 0.05, 0.6, [f"**{it.get('label', '')}**", it.get("photo", "")], 11, self.fg(s._bg), after=0.1, check=False)
                cx += cw
            if ad:
                self.rect(s, cx, y, wa, 1.0, "negro")
                self.tb(s, cx + 0.12, y + 0.08, wa - 0.2, 0.6, ad.get("code", ""), 30, self.accent("negro"), "display", check=False)
                self.tb(s, cx, y + 1.1, wa, 0.6, [f"**{ad.get('label', '')}**", ad.get("photo", "")], 11, self.fg(s._bg), after=0.1, check=False)
            if b.get("legend"):
                self.tb(s, x, y + 1.8, w_, 0.3, "   ".join(f"■ {k}: {v}" for k, v in b["legend"]), 10, self.fg(s._bg), "mono", check=False)
        return self.framed({"caption": b.get("caption")}, w, 2.15, d)

    def u_compare(self, b, w):
        L, R = b.get("left") or {}, b.get("right") or {}
        hi = max(len(L.get("items", [])), len(R.get("items", [])))
        h = 1.0 + hi * 0.5 + (0.75 if b.get("verdict") else 0)

        def d(s, x, y, w_, h_):
            hv = 0.75 if b.get("verdict") else 0
            cw = (w_ - GAP) / 2
            for k, side in enumerate((L, R)):
                bg = "crema" if k == 0 else "bosque"
                cx = x + k * (cw + GAP)
                self.rect(s, cx, y, cw, h_ - hv, bg, lw=3)
                self.tb(s, cx + 0.25, y + 0.15, cw - 0.5, 0.5, side.get("title", ""), 26, self.fg(bg, ("negro", "crema")), "display", upper=True, check=False)
                if side.get("sub"):
                    self.tb(s, cx + 0.25, y + 0.65, cw - 0.5, 0.3, side["sub"], 11, self.fg(bg, ("negro", "crema")), "mono", upper=True, check=False)
                self.bullets(s, cx + 0.25, y + 1.0, cw - 0.5, h_ - hv - 1.1, side.get("items", []), 14, self.fg(bg, ("negro", "crema")))
            if hv:
                self.rect(s, x, y + h_ - hv + 0.12, w_, hv - 0.12, "negro", lw=0)
                self.tb(s, x + 0.25, y + h_ - hv + 0.2, w_ - 0.5, hv - 0.25, b["verdict"], 15, self.accent("negro"), bold=True, anchor=MSO_ANCHOR.MIDDLE)
        return [(h, d)]

    def u_image(self, b, w):
        its = b.get("items") or [b]
        r = b.get("ratio", "16/9")
        try:
            a, c = (float(v) for v in str(r).split("/"))
        except ValueError:
            a, c = 16, 9
        hb = min(4.3, (w / len(its)) * c / a)

        def d(s, x, y, w_, h_):
            cw = (w_ - GAP * (len(its) - 1)) / len(its)
            for k, it in enumerate(its):
                self.picture(s, it.get("src"), x + k * (cw + GAP), y, cw, h_, it.get("treatment") or b.get("treatment"), it.get("need"), it.get("caption") or it.get("alt"))
        cap = b.get("caption") if not b.get("items") else " · ".join(plain(i.get("caption", "")) for i in its if i.get("caption"))
        return self.framed({"title": b.get("title"), "caption": cap}, w, hb, d)

    def u_media(self, b, w):
        def d(s, x, y, w_, h_):
            f = self.img_root / b.get("src", "")
            if not f.is_file():
                self.rep["avisos"].append({"bloque": b.get("src", "media"), "causa": "archivo de medio no encontrado", "ajuste": "se dibujó el hueco"})
                return self.picture(s, None, x, y, w_, h_, need=f"{b.get('kind', 'video')}: {b.get('src', '')}")
            mb = f.stat().st_size / 1e6
            if mb > 50:
                self.rep["avisos"].append({"bloque": b["src"], "causa": f"medio de {mb:.0f} MB", "ajuste": "el PPTX pesará mucho; recorta o comprime"})
            if b.get("kind") == "gif":
                pic = s.shapes.add_picture(str(f), Inches(x), Inches(y), height=Inches(h_))   # el GIF anima en modo presentación
                pic.left = int(Inches(x) + (Inches(w_) - pic.width) / 2)
            else:
                poster = self.img_root / b["poster"] if b.get("poster") and (self.img_root / b["poster"]).is_file() else None
                s.shapes.add_movie(str(f), Inches(x), Inches(y), Inches(w_), Inches(h_), poster_frame_image=str(poster) if poster else None, mime_type="video/mp4")
        cap = " · ".join(plain(v) for v in (b.get("caption"), b.get("credit")) if v)
        return self.framed({"caption": cap}, w, 4.0, d)

    def u_toc(self, b, w):
        def d(s, x, y, w_, h_):
            secs = [(sec["n"], plain(sec["title"])) for p in self.parts for sec in p["sections"] if sec.get("toc", True) and not p.get("_aux")]
            half = math.ceil(len(secs) / 2)
            for k, col in enumerate((secs[:half], secs[half:])):
                self.tb(s, x + k * w_ / 2, y, w_ / 2 - GAP, h_, [f"**{n}**  {t}" for n, t in col], 14, self.fg(s._bg), lh=1.1, after=0.25, check=False)
        return [(BOT - TOP_CONT - 0.2, d)]

    def u_flowchart(self, b, w):
        nodes, edges = {n["id"]: n for n in b.get("nodes", [])}, b.get("edges", [])
        layer = {k: 0 for k in nodes}
        for _ in range(len(nodes)):                  # capa = camino más largo desde un inicio (sin ciclos: tope de iteraciones)
            for e in edges:
                if e.get("from") in layer and e.get("to") in layer and layer[e["to"]] < layer[e["from"]] + 1 <= len(nodes):
                    layer[e["to"]] = layer[e["from"]] + 1
        cols = {}
        for k, v in layer.items():
            cols.setdefault(v, []).append(k)
        shp = {"start": MSO_SHAPE.FLOWCHART_TERMINATOR, "end": MSO_SHAPE.FLOWCHART_TERMINATOR, "decision": MSO_SHAPE.DIAMOND, "io": MSO_SHAPE.PARALLELOGRAM}

        def d(s, x, y, w_, h_):
            nl = max(cols) + 1 if cols else 1
            cw = w_ / nl
            made = {}
            for c, ids in cols.items():
                rh = h_ / len(ids)
                for r, nid in enumerate(ids):
                    n = nodes[nid]
                    kind = n.get("kind", "process")
                    bw, bh = min(cw - 0.5, 2.3), min(rh - 0.35, 1.1 if kind == "decision" else 0.8)
                    bg = "negro" if kind in ("start", "end") else ("ocre" if kind == "decision" else "crema")
                    sh = self.rect(s, x + c * cw + (cw - bw) / 2, y + r * rh + (rh - bh) / 2, bw, bh, bg, shape=shp.get(kind, MSO_SHAPE.RECTANGLE))
                    tf = sh.text_frame
                    tf.word_wrap = True
                    tf.margin_left = tf.margin_right = Inches(0.05)
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    self.runs(p, n.get("label", ""), 12, self.fg(bg, ("crema", "negro")) if bg == "negro" else self.fg(bg), bold=kind in ("start", "end"))
                    made[nid] = sh
            for e in edges:
                a, z = made.get(e.get("from")), made.get(e.get("to"))
                if not (a and z):
                    continue
                cx = s.shapes.add_connector(MSO_CONNECTOR.ELBOW, 0, 0, 0, 0)
                fwd = layer[e["to"]] > layer[e["from"]]
                cx.begin_connect(a, 3 if fwd else 2)
                cx.end_connect(z, 1 if fwd else 0)
                cx.line.color.rgb, cx.line.width = self.C[self.fg(s._bg)], Pt(2)
                etree.SubElement(cx.line._get_or_add_ln(), qn("a:tailEnd"), type="triangle", w="med", len="med")
                if e.get("label"):
                    self.tb(s, (a.left + z.left + a.width) / 2 / 914400 - 0.4, (a.top + z.top) / 2 / 914400 - 0.05, 0.8, 0.25, e["label"], 10, self.fg(s._bg), "mono", True, align=PP_ALIGN.CENTER, upper=True, check=False)
        return self.framed(b, w, 4.0, d)

    def u_custom(self, b, w):
        self.rep["avisos"].append({"bloque": "custom", "causa": "HTML libre no se traduce a PowerPoint", "ajuste": "se mostró su texto como nota"})
        return self.u_callout({"tone": "nota", "text": re.sub(r"<[^>]+>", " ", b.get("html", ""))[:500]}, w)

    # ---------- diapositivas de estructura ----------
    def cover(self, cov, meta):
        s = self.slide("negro")
        ink, acc = self.fg("negro", ("crema",)), self.accent("negro")
        img = cov.get("image") and self.picture(s, cov["image"], W * 0.56, 0, W * 0.44, H, cov.get("treatment") or "ocre", caption=cov.get("alt"), decor=True)
        v = s.shapes.add_textbox(Inches(-2.75), Inches(2.9), Inches(7.4), Inches(1.6))     # palabra vertical: sangra por el borde izquierdo
        v.rotation = 270
        v.name = "sangrado"
        tf = v.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        word = str(cov.get("vertical", "")).upper()
        self.runs(tf.paragraphs[0], word, min(120, int(7.2 * 72 / (max(len(word), 1) * CHAR_W["display"] * 1.05))), acc, "display")
        xt = 2.35
        wt = (W * 0.56 if img else W - M) - xt - 0.3
        hl = meta.get("header_left") or []
        self.tb(s, xt, 0.5, wt, 0.6, hl[:2], 11, ink, "mono", True, upper=True, after=0.1, check=False)
        self.rect(s, xt, 3.55, wt, 0, line=ink, lw=6)
        tt = plain(cov.get("title", "")).upper()
        self.tb(s, xt, 3.75, wt, 1.6, cov.get("title", ""), fit(tt, wt, 1.6, 44, 22, 0.95, "display"), ink, "display", upper=True, lh=0.95)
        self.tb(s, xt, 5.45, wt, 1.15, cov.get("text", ""), fit(plain(cov.get("text", "")), wt, 1.15, 17, 11, 1.15), ink)
        self.tb(s, xt, 6.75, wt, 0.3, "   ·   ".join(v for v in (meta.get("date"), meta.get("version")) if v), 11, ink, "mono", True, upper=True, check=False)
        self.notes(s, plain(cov.get("text", "")))

    def agenda(self, parts):
        s = self.slide("arena")
        self.chrome(s, "Agenda")
        ink = self.fg("arena")
        self.tb(s, M, 1.0, CW, 1.0, "Agenda", 60, ink, "display", upper=True, check=False)
        n = len(parts)
        cw = (CW - GAP * (n - 1)) / max(n, 1)
        for k, p in enumerate(parts):
            cx = M + k * (cw + GAP)
            self.rect(s, cx, 2.4, cw, 4.1, "crema" if k % 2 else "negro", lw=3)
            bg = "crema" if k % 2 else "negro"
            self.tb(s, cx + 0.2, 2.5, cw - 0.4, 1.3, f"{k+1:02d}", 64, self.accent(bg) if bg == "negro" else "hoja" if ("hoja", "crema") in self.pairs else self.fg(bg), "display", check=False)
            self.tb(s, cx + 0.2, 3.85, cw - 0.4, 1.0, p["title"], fit(plain(p["title"]).upper(), cw - 0.4, 1.0, 20, 11), self.fg(bg, ("crema", "negro")), bold=True, upper=True)
            self.tb(s, cx + 0.2, 4.9, cw - 0.4, 1.5, p.get("desc", ""), fit(plain(p.get("desc", "")), cw - 0.4, 1.5, 12, 9.5, 1.15), self.fg(bg, ("crema", "negro")))

    def divider(self, k, part):
        bg = part.get("bg") if part.get("bg") in self.C else ("bosque" if k % 2 else "negro")
        s = self.slide(bg)
        ink = self.fg(bg, ("crema",))
        num = part.get("num_color") if (part.get("num_color"), bg) in self.pairs else self.accent(bg)
        if part.get("image"):
            self.picture(s, part["image"], W * 0.6, 0, W * 0.4, H, "negroverde", decor=True)
        self.tb(s, M - 0.1, -0.35, 7.5, 4.6, f"{k+1:02d}", 330, num, "display", lh=0.85, check=False).name = "sangrado"   # el número se recorta por arriba
        self.rect(s, M, 4.55, W * 0.5, 0, line=ink, lw=6)
        self.tb(s, M, 4.7, W * 0.55, 1.2, part["title"], fit(plain(part["title"]).upper(), W * 0.55, 1.2, 54, 28, 0.95, "display"), ink, "display", upper=True, lh=0.95)
        self.tb(s, M, 5.95, W * 0.5, 0.9, part.get("desc", ""), fit(plain(part.get("desc", "")), W * 0.5, 0.9, 15, 11, 1.15), ink)
        self.tb(s, M, 6.95, 4, 0.3, f"Parte {k+1:02d} / {len(self.parts):02d}", 10, ink, "mono", True, upper=True, check=False)
        self.notes(s, plain(part.get("desc", "")))

    def section(self, part, sec):
        crumb = f"{sec['n']} · {part['title']}"
        if sec.get("tema") and plain(sec["tema"]).lower() not in plain(sec["title"]).lower():
            crumb += f" · {plain(sec['tema'])}"
        s = self.slide("crema")
        self.chrome(s, crumb)
        ink = self.fg("crema")
        ts = fit(plain(sec["title"]).upper(), CW, 1.25, 38, 24, 0.95, "display")
        th = est_h(plain(sec["title"]).upper(), CW, ts, 0.95, font="display") + 0.1
        self.tb(s, M, 0.98, CW, th, sec["title"], ts, ink, "display", upper=True, lh=0.95)
        y = 0.98 + th + 0.15
        if sec.get("lead"):
            if sec.get("kind") == "clave":
                lh_ = est_h(plain(sec["lead"]), CW - 0.6, 20, 1.15, 0.35) + 0.4
                self.rect(s, M, y, CW, lh_, "bosque", lw=0)
                self.tb(s, M + 0.3, y + 0.2, CW - 0.6, lh_ - 0.3, sec["lead"], 20, self.fg("bosque", ("crema",)), bold=True)
            else:
                lh_ = est_h(plain(sec["lead"]), CW, 20, 1.15, 0.35)
                self.tb(s, M, y, CW, lh_, sec["lead"], 20, ink, bold=True)
            y += lh_ + GAP
        self.notes(s, "\n".join(f"[{n_.get('n')}] {plain(n_.get('text', ''))}" for n_ in sec.get("notes", [])))
        first_y = y
        for h, draw in self.units(sec.get("blocks", [])):
            # nueva diapositiva si ya hay algo en esta, o si en una limpia cabría y aquí no
            if y + h > BOT and (y > first_y + 0.01 or (h <= BOT - TOP_CONT and y > TOP_CONT + 0.01)):
                self._density(s)
                s = self.slide("crema")
                self.chrome(s, crumb + " (cont.)")
                y = first_y = TOP_CONT
            hh = min(h, BOT - y)
            if hh < h - 0.05:
                self.rep["avisos"].append({"bloque": f"{sec['id']} · diapositiva {self.n}", "causa": f"bloque de {h:.1f} in en {hh:.1f} in", "ajuste": "se redujo; divide el bloque si se ve apretado"})
            draw(s, M, y, CW, hh)
            y += hh + GAP
        self._density(s)

    def _density(self, s):
        txt = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if len(txt.split()) > 170:
            self.rep["avisos"].append({"bloque": f"diapositiva {self.n}", "causa": f"{len(txt.split())} palabras", "ajuste": "muy densa para proyectar; reparte el texto o pasa detalles a las notas"})

    def credits_slide(self):
        if not self.used_credits:
            return
        s = self.slide("crema")
        self.chrome(s, "Créditos")
        self.tb(s, M, 1.0, CW, 0.9, "Créditos de imágenes", 40, self.fg("crema"), "display", upper=True, check=False)
        self.bullets(s, M, 2.1, CW, 4.5, [f"{c.get('especie_o_tema', k)} — {c.get('autor', '')} · {c.get('licencia', '')} · {c.get('url_commons', '')}" for k, c in self.used_credits.items()], 12, self.fg("crema"))

    # ---------- auditoría de la presentación ----------
    def audit(self):
        out, small = [], []
        for i, s in enumerate(self.p.slides, 1):
            for sh in s.shapes:
                if sh.name == "sangrado" or sh.width is None:
                    continue
                if sh.left < -Emu(1000) or sh.top < -Emu(1000) or sh.left + sh.width > self.p.slide_width + Emu(9000) or sh.top + sh.height > self.p.slide_height + Emu(9000):
                    out.append({"diapositiva": i, "forma": sh.name})
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            if r.font.size and r.font.size.pt < 9.5 and r.text.strip():
                                small.append({"diapositiva": i, "texto": r.text[:30], "pt": r.font.size.pt})
        a = self.rep["auditoria"]
        a["fuera_de_diapositiva"], a["letra_pequena"] = out, small[:20]


def run(doc_path, out_pptx, images_root):
    doc = json.loads(Path(doc_path).read_text(encoding="utf-8"))
    design_dir, m = render.load_design(doc.get("design", "anura"))
    warns = []
    doc = {**doc, "format": "pptx-16x9"} if doc.get("format") in (None, "pdf-a4") else doc     # el mismo JSON del PDF sirve para PPTX
    errs = contract.validate(doc, {**m, "formats": list(dict.fromkeys(m["formats"] + ["pptx-16x9"]))}, warns)
    rep = {"estado": "ok", "formato": "pptx-16x9", "errores": errs, "avisos": warns, "auditoria": {}, "imagenes_necesarias": []}
    if errs:
        rep["estado"] = "rechazado"
        return rep
    roles = render.resolve_roles(doc, design_dir, m, rep["avisos"])
    rep["avisos"] += paleta.check(roles, m)
    dk = Deck(roles, m, images_root, rep)
    meta, d = doc.get("meta", {}), doc["document"]
    dk.footer = meta.get("footer") or meta.get("title", "")
    dk.parts = [dict(p) for p in d.get("parts", [])]
    n = 0
    for p in dk.parts:
        p["sections"] = [dict(s) for s in p.get("sections", [])]
        for sec in p["sections"]:
            n += 1
            sec["n"] = f"{n:02d}"
    dk.cover(d["cover"], meta)
    if len(dk.parts) > 1:
        dk.agenda(dk.parts)
    for k, part in enumerate(dk.parts):
        if len(dk.parts) > 1:
            dk.divider(k, part)
        for sec in part["sections"]:
            dk.section(part, sec)
    dk.credits_slide()
    dk.audit()
    cp = dk.p.core_properties
    cp.title, cp.author, cp.subject = meta.get("title", ""), meta.get("author", ""), meta.get("subject", "")
    cp.keywords = ", ".join(meta.get("keywords", []))
    Path(out_pptx).parent.mkdir(parents=True, exist_ok=True)
    dk.p.save(out_pptx)
    rep["diapositivas"] = len(dk.p.slides)
    rep["fuentes"] = {"usa": sorted(set(dk.F.values())), "reemplazo_si_no_estan": {k: v for k, v in FONT_SWAP.items() if k in dk.F.values()},
                      "nota": "PowerPoint no incrusta estas fuentes: instala los .ttf de /v1/designs/<id>/fonts.zip (tools/fuentes_ttf.py) o se verán con el reemplazo"}
    a = rep["auditoria"]
    if a.get("desbordes") or a.get("fuera_de_diapositiva"):
        rep["estado"] = "con_avisos"
    return rep


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("doc")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--imagenes", default="imagenes")
    a = ap.parse_args()
    rep = run(a.doc, a.out, a.imagenes)
    Path(a.out).with_suffix(".informe.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: rep[k] for k in ("estado", "diapositivas") if k in rep} | {"avisos": len(rep["avisos"]), "desbordes": len(rep["auditoria"].get("desbordes", []))}, ensure_ascii=False))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")      # Windows: la consola cp1252 no acepta ✔ → ·
    main()
