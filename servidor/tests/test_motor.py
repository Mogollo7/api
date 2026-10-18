# [CLAUDE-CONSTRUCTOR] Pruebas del motor creadas por Claude. Ejecutar:  python -m unittest discover -s servidor/tests -v
import json
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE.parent / "tools"))

import render  # noqa: E402
from engine import blocks, charts, contract, diagram, muestras, paleta, prompt, redundancia, spec, text, tpl, variants  # noqa: E402

DD = HERE.parent / "designs/anura"
M = json.loads((DD / "manifest.json").read_text(encoding="utf-8"))
EJ = HERE / "ejemplo_minimo.json"
IMG = HERE.parent.parent / "teoria/imagenes"


class Texto(unittest.TestCase):
    def test_marcas(self):
        h = text.inline("**a** *b* `c` [FUENTE] dato pendiente <script>")
        for s in ("<b>a</b>", "<em>b</em>", "<code>c</code>", 'class="et f"', 'class="pend"'):
            self.assertIn(s, h)
        self.assertNotIn("<script>", h)

    def test_coma_decimal(self):
        self.assertEqual(text.num(62.85), "62,85")


class Contrato(unittest.TestCase):
    def doc(self):
        return json.loads(EJ.read_text(encoding="utf-8"))

    def test_valido(self):
        w = []
        self.assertEqual(contract.validate(self.doc(), M, w), [])
        self.assertEqual(w, [])

    def test_no_estandar_es_aviso_no_error(self):
        d, w = self.doc(), []
        d["document"]["parts"][0]["sections"][0]["blocks"].append({"type": "html", "text": "<b>x</b>"})
        self.assertEqual(contract.validate(d, M, w), [])
        self.assertTrue(any("vocabulario" in x["causa"] for x in w))

    def test_sin_portada_se_rechaza(self):
        d = self.doc()
        del d["document"]["cover"]
        self.assertTrue(contract.validate(d, M, []))


class Graficos(unittest.TestCase):
    R = M["palette"]["roles"]

    def test_todos_los_tipos_dibujan(self):
        datos = {
            "bar": {"items": [{"l": "a", "v": 1}]}, "hbar": {"items": [{"l": "a", "v": 1}]}, "line": {"points": [[1, 1], [2, 2]]},
            "stack": {"rows": [{"l": "a", "v": [1, 2, 3]}]}, "hist": {"bars": [[1, 2], [2, 3]]},
            "box": {"groups": [{"l": "a", "lo": 1, "q1": 2, "med": 3, "q3": 4, "hi": 5, "out": [9]}]}, "scatter": {"points": [[1, 1], [2, 3], [3, 2]], "trend": "auto"},
            "heat": {"labels": ["a", "b"], "matrix": [[1, .2], [.2, 1]]},
            "map": {"outline": "colombia", "points": [{"lon": -74, "lat": 4.7}]},
        }
        out = {}
        for t, d in datos.items():
            out[t] = charts.draw({"type": t, "data": d}, self.R, {"colombia": json.loads((DD / "geo/colombia.json").read_text(encoding="utf-8"))["poligonos"]})
            self.assertTrue(out[t].startswith("<svg"), t)

    def test_limites(self):
        with self.assertRaises(charts.ChartError):
            charts.draw({"type": "bar", "data": {"items": [{"l": str(i), "v": i} for i in range(30)]}}, self.R)

    def test_flowchart_ciclo_y_limite(self):
        n = [{"id": "a", "label": "A"}, {"id": "b", "label": "B", "kind": "decision"}]
        self.assertIn("polygon", diagram.draw({"nodes": n, "edges": [{"from": "a", "to": "b"}, {"from": "b", "to": "a", "label": "no"}]}))
        with self.assertRaises(diagram.DiagramError):
            diagram.draw({"nodes": [{"id": str(i), "label": "x"} for i in range(15)], "edges": []})


class Paleta(unittest.TestCase):
    def test_anura_y_cine_sin_avisos_de_contraste(self):
        pals = json.loads((DD / "palettes.json").read_text(encoding="utf-8"))
        for k in ("anura", "cine"):
            r = dict(M["palette"]["roles"])
            r.update(pals[k])
            self.assertEqual([x for x in paleta.check(r, M) if "contraste" in x["causa"]], [], k)

    def test_detecta_gris_confundible(self):
        r = dict(M["palette"]["roles"])
        r.update({"hoja": "#7a7a7a", "vivo": "#7c7c7c"})
        self.assertTrue(any("hoja / vivo" in x["bloque"] for x in paleta.check(r, M)))


class Variantes(unittest.TestCase):
    def test_resolver(self):
        reg = variants.Registry(DD)
        w = []
        self.assertEqual(reg.resolve("portada", "auto", w), "CV-1")
        self.assertEqual(reg.resolve("portada", "CV-9", w), "CV-1")
        self.assertTrue(w)
        reg.set_state("DV-E", "rechazada")
        self.assertEqual(reg.resolve("divisor", "DV-E", w), "DV-A")


class Render(unittest.TestCase):
    def run_doc(self, p):
        with tempfile.TemporaryDirectory() as t:
            rep = render.run(p, Path(t) / "x.pdf", IMG)
            html = (Path(t) / "x.html").read_text(encoding="utf-8") if (Path(t) / "x.html").exists() else ""
            return rep, html

    def test_metadatos_y_etiquetado(self):
        import subprocess
        with tempfile.TemporaryDirectory() as t:
            render.run(EJ, Path(t) / "x.pdf", IMG)
            info = subprocess.run(["pdfinfo", "-enc", "UTF-8", str(Path(t) / "x.pdf")], capture_output=True, text=True, encoding="utf-8", errors="replace").stdout
        self.assertIn("Ejemplo mínimo", info)
        self.assertIn("Tagged:          yes", info)

    def test_minimo_sin_problemas(self):
        rep, _ = self.run_doc(EJ)
        self.assertEqual(rep["errores"], [])
        self.assertEqual(rep["paginas"], 5)
        self.assertEqual(rep["auditoria"], {})

    def test_no_estandar_no_rompe(self):
        rep, html = self.run_doc(HERE / "no_estandar.json")
        self.assertEqual(rep["errores"], [])
        self.assertGreater(rep["paginas"], 3)
        self.assertEqual(rep["auditoria"], {})
        self.assertGreaterEqual(len(rep["avisos"]), 10)
        self.assertNotIn("<script>alert", html)
        self.assertNotIn("onclick", html)

    def test_correccion_de_contraste(self):
        d = json.loads(EJ.read_text(encoding="utf-8"))
        d["palette"] = {"hoja": "#bdbdbd"}   # texto en 'hoja' sobre papel: no se lee
        with tempfile.TemporaryDirectory() as t:
            p = Path(t) / "d.json"
            p.write_text(json.dumps(d))
            rep = render.run(p, Path(t) / "x.pdf", IMG)
        self.assertTrue(rep["correcciones_contraste"])
        for c in rep["correcciones_contraste"]:
            self.assertGreaterEqual(c["ratio_despues"], 3)


class Redundancia(unittest.TestCase):
    def test_similar(self):
        self.assertTrue(redundancia.similar("Glosario, fuentes y pendientes", "Fuentes y pendientes"))
        self.assertTrue(redundancia.similar("Glosario, fuentes", "Glosario", laxo=True))
        self.assertFalse(redundancia.similar("Enfermedad en fauna y flora", "Fauna"))   # subtítulo legítimo de una palabra
        self.assertFalse(redundancia.similar("Histograma de las calificaciones", "Nota vs. valoración pública"))

    def test_el_servidor_omite_lo_repetido(self):
        d = json.loads(EJ.read_text(encoding="utf-8"))
        sec = d["document"]["parts"][0]["sections"][1]
        sec["title"], sec["tema"] = "Gráficos y tablas de muestra", "Gráficos y tablas"
        sec["blocks"].insert(0, {"type": "h", "text": "Gráficos y tablas"})
        with tempfile.TemporaryDirectory() as t:
            p = Path(t) / "d.json"
            p.write_text(json.dumps(d))
            rep = render.run(p, Path(t) / "x.pdf", IMG)
        self.assertTrue(any("redundancia" in a["ajuste"] for a in rep["avisos"]))
        self.assertEqual([x for x in rep["auditoria"].get("794-print", []) if x.get("tipo") == "redundancia"], [])


class Plantillas(unittest.TestCase):
    def test_tpl(self):
        self.assertEqual(tpl.render("{{a}}|{{&b}}|{{#if c}}Sí{{else}}No{{/if}}|{{#each l}}[{{x}}]{{/each}}", {"a": "<", "b": "<", "c": 0, "l": [1, {"x": "z"}], "x": "o"}), "&lt;|<|No|[o][z]")

    def test_cada_diseno_trae_plantillas(self):
        for d in ("anura", "editorial"):
            self.assertTrue(list((HERE.parent / "designs" / d / "templates").glob("cover.*.html")), d)
            self.assertTrue(list((HERE.parent / "designs" / d / "templates").glob("divider.*.html")), d)


class Spec(unittest.TestCase):
    def test_schema_y_ejemplos(self):
        sc = spec.json_schema()
        self.assertEqual(sc["properties"]["contract"]["const"], spec.CONTRACT_VERSION)
        for name, b in spec.BLOCKS.items():
            w = []
            spec.validate_block(b["example"], name, w)
            self.assertEqual(w, [], name)

    def test_prompt_maestro(self):
        for d in ("anura", "editorial"):
            for corto in (False, True):
                p = prompt.prompt_maestro(d, corto=corto)
                for k in ("agente de contenido", "Anti-redundancia", "dato pendiente", "{{CONTENIDO}}", "flowchart"):
                    self.assertIn(k, p)
        self.assertLess(len(prompt.prompt_maestro("anura", corto=True)), len(prompt.prompt_maestro("anura")))

    def test_muestra_valida(self):
        d = muestras.doc_muestra("anura")
        self.assertEqual(contract.validate(d, M, []), [])


class BloquesNuevos(unittest.TestCase):
    R = M["palette"]["roles"]

    def ctx(self):
        return blocks.Ctx(M, self.R, IMG, {}, HERE.parent / ".cache", [], DD)

    def test_cita_codigo_refs_sublista(self):
        c = self.ctx()
        self.assertIn("<blockquote", blocks.render({"type": "quote", "text": "x", "by": "y"}, c))
        h = blocks.render({"type": "code", "code": "a<b", "lang": "py"}, c)
        self.assertIn("a&lt;b", h)
        self.assertIn("ul class=\"bl sub\"", blocks.render({"type": "list", "items": [{"t": "a", "items": ["b"]}]}, c))
        self.assertIn('class="fn"', text.inline("dato[^1]"))

    def test_graficos_nuevos(self):
        d = {"pie": {"items": [{"l": "a", "v": 1}, {"l": "b", "v": 3}]}, "bar": {"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}, {"name": "t", "values": [2, 1]}], "mode": "stacked"},
             "line": {"series": [{"name": "s", "points": [["2026-01", 1], ["2026-02", 3]]}]}}
        for t, data in d.items():
            self.assertTrue(charts.draw({"type": t, "data": data}, self.R).startswith("<svg"), t)


class Diseno(unittest.TestCase):
    def test_validar_editorial_rapido(self):
        import subprocess
        r = subprocess.run([sys.executable, str(HERE.parent / "validar_diseno.py"), "editorial", "--rapido"], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stdout[-800:])


class Comons(unittest.TestCase):
    def test_parse_sin_red(self):
        import commons
        resp = {"query": {"pages": {"1": {"title": "File:Rana.jpg", "imageinfo": [{"thumburl": "https://x/y.jpg", "descriptionurl": "https://c/Rana", "extmetadata": {
            "LicenseShortName": {"value": "CC BY-SA 4.0"}, "Artist": {"value": "<a>Ana</a>"}, "ImageDescription": {"value": "Una rana"}}}]},
            "2": {"title": "File:No.jpg", "imageinfo": [{"thumburl": "https://x/n.jpg", "extmetadata": {"LicenseShortName": {"value": "All rights reserved"}}}]}}}}
        r = commons.parse(resp)
        self.assertEqual(len(r), 1)
        self.assertEqual(r[0]["autor"], "Ana")


if __name__ == "__main__":
    unittest.main()


class Pptx(unittest.TestCase):
    """El mismo JSON del PDF da un PPTX válido, con objetos nativos y sin formas fuera de la diapositiva; lo no estándar se degrada."""
    def test_ejemplo_y_no_estandar(self):
        import render_pptx
        from pptx import Presentation
        for src in (EJ, HERE / "no_estandar.json"):
            with tempfile.TemporaryDirectory() as t:
                out = Path(t) / "x.pptx"
                rep = render_pptx.run(src, out, IMG)
                self.assertNotEqual(rep["estado"], "rechazado", rep["errores"])
                self.assertEqual(rep["auditoria"]["fuera_de_diapositiva"], [])
                p = Presentation(out)
                self.assertEqual(len(p.slides), rep["diapositivas"])
                kinds = {sh.shape_type for s in p.slides for sh in s.shapes}
                if src == EJ:
                    from pptx.enum.shapes import MSO_SHAPE_TYPE
                    self.assertIn(MSO_SHAPE_TYPE.CHART, kinds)      # gráfico nativo editable, no imagen
                    self.assertIn(MSO_SHAPE_TYPE.TABLE, kinds)

    def test_texto_solo_en_pares_permitidos(self):
        import render_pptx
        from engine import paleta as pal
        dk = render_pptx.Deck(M["palette"]["roles"], M, IMG, {"avisos": [], "auditoria": {}, "imagenes_necesarias": []})
        for bg in M["palette"]["roles"]:
            self.assertGreaterEqual(pal.ratio(dk.roles[dk.fg(bg)], dk.roles[bg]), 4.5, bg)


class Video(unittest.TestCase):
    """El director arma una composición HyperFrames válida desde el mismo JSON, con el texto escapado (es dato, no HTML)."""
    def test_composicion(self):
        import render_video
        doc = json.loads(EJ.read_text(encoding="utf-8"))
        doc["document"]["cover"]["title"] = "<script>alert(1)</script> Título"
        with tempfile.TemporaryDirectory() as t:
            rep = {"avisos": []}
            dr, f = render_video.build(doc, Path(t), "video-9x16", rep)
            h = f.read_text(encoding="utf-8")
            self.assertNotIn("<script>alert", h)
            self.assertIn('data-composition-id="main"', h)
            self.assertIn('window.__timelines["main"]=tl', h)
            self.assertIn(f'data-duration="{dr.t:.3f}"', h)
            self.assertGreater(len(dr.scenes), 2)
            self.assertTrue((Path(t) / "assets" / "gsap.min.js").exists())

    def test_cifra(self):
        import render_video
        self.assertEqual(render_video.number("$ 128 M"), ("$ ", 128.0, " M", 0))
        self.assertEqual(render_video.number("62,9 %"), ("", 62.9, " %", 1))
        self.assertIsNone(render_video.number("sin dato"))


class Oficina(unittest.TestCase):
    """Word y Excel desde el mismo contrato; Excel con números reales (coma decimal) y totales con fórmula."""
    def test_numeros(self):
        import render_office as ro
        self.assertEqual(ro.num("1.234,5"), 1234.5)
        self.assertEqual(ro.num("62,9 %"), 62.9)
        self.assertEqual(ro.num(12), 12.0)
        self.assertIsNone(ro.num("dato pendiente"))

    def test_docx_y_xlsx(self):
        import render_office as ro
        with tempfile.TemporaryDirectory() as t:
            for ext in ("docx", "xlsx"):
                rep = ro.run(EJ, Path(t) / f"x.{ext}", IMG)
                self.assertEqual(rep["estado"], "ok", rep["errores"])
                self.assertGreater((Path(t) / f"x.{ext}").stat().st_size, 3000)


class Api(unittest.TestCase):
    def test_parches_y_firma(self):
        import api
        doc = json.loads(EJ.read_text(encoding="utf-8"))
        sid = doc["document"]["parts"][0]["sections"][0]["id"]
        api.apply_patch(doc, [{"op": "insert", "section": sid, "block": 0, "value": {"type": "p", "text": "nuevo"}}])
        self.assertEqual(doc["document"]["parts"][0]["sections"][0]["blocks"][0]["text"], "nuevo")
        with self.assertRaises(ValueError):
            api.apply_patch(doc, [{"op": "set", "section": "no-existe", "block": 0, "value": {}}])
        url = api.sign("/v1/files/abc.pdf")
        path, q = url.split("?")
        qs = {k: [v] for k, v in (x.split("=") for x in q.split("&"))}
        self.assertTrue(api.signed_ok(path, qs))
        self.assertFalse(api.signed_ok("/v1/files/otro.pdf", qs))           # la firma no sirve para otro archivo
